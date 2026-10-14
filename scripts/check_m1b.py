"""M1b smoke test inside the API container: upload -> index -> Qdrant.

Uses synthetic fixtures only. First indexing may download the ~90MB CPU
embedding model; later runs reuse the persistent cache volume.
"""

import asyncio
from io import BytesIO
from pathlib import Path
from uuid import UUID

import asyncpg
import httpx
import pymupdf
from docx import Document
from qdrant_client import AsyncQdrantClient

from app.core.config import get_settings
from app.storage.qdrant import DOCUMENTS_COLLECTION, _filter_document


def make_pdf() -> bytes:
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "Synthetic contract: payment is due within 30 days.")
    raw = pdf.tobytes()
    pdf.close()
    return raw


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Payment Clause", level=1)
    document.add_paragraph("Payment is due within 30 days after acceptance.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Party A"
    table.cell(0, 1).text = "Example Corporation"
    memory = BytesIO()
    document.save(memory)
    return memory.getvalue()


async def assert_indexed(document_id: str, expected_count: int, filename: str) -> None:
    settings = get_settings()
    client = AsyncQdrantClient(url=settings.qdrant_url, timeout=30)
    try:
        points, _ = await client.scroll(
            collection_name=DOCUMENTS_COLLECTION,
            scroll_filter=_filter_document(document_id),
            with_payload=True,
            with_vectors=True,
            limit=100,
        )
        assert len(points) == expected_count, (filename, len(points), expected_count)
        for point in points:
            payload = point.payload or {}
            assert payload["document_id"] == document_id
            assert payload["filename"] == filename
            assert payload["chunk_id"]
            assert payload["text"]
            assert isinstance(point.vector, list)
            assert len(point.vector) == 512

        if filename.endswith(".pdf"):
            assert all(point.payload["page"] == 1 for point in points)
        elif filename.endswith(".docx"):
            assert all(point.payload["section"] == "Payment Clause" for point in points)
            assert all(point.payload["page"] is None for point in points)

    finally:
        await client.close()

    conn = await asyncpg.connect(settings.postgres_dsn)
    try:
        row = await conn.fetchrow(
            "SELECT status, chunk_count, embedding_model FROM documents WHERE id=$1",
            UUID(document_id),
        )
        assert row["status"] == "indexed"
        assert row["chunk_count"] == expected_count
        assert row["embedding_model"] == settings.embedding_model
    finally:
        await conn.close()


async def main() -> None:
    example = Path("data/samples/sample_contract.txt").read_bytes()
    cases = [
        ("sample_contract.txt", example, "text/plain"),
        ("synthetic.pdf", make_pdf(), "application/pdf"),
        (
            "synthetic.docx",
            make_docx(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
    ]

    async with httpx.AsyncClient(
        base_url="http://127.0.0.1:8000",
        timeout=600,
        trust_env=False,
    ) as client:
        for index, (filename, data, mime) in enumerate(cases):
            upload = await client.post(
                "/v1/documents",
                files={"file": (filename, data, mime)},
            )
            upload.raise_for_status()
            result = upload.json()
            doc_id = result["document_id"]
            assert result["status"] == "uploaded"

            response = await client.post(f"/v1/documents/{doc_id}/index")
            response.raise_for_status()
            indexed = response.json()
            assert indexed["status"] == "indexed"
            assert indexed["vector_size"] == 512
            assert indexed["chunk_count"] >= 1
            await assert_indexed(doc_id, indexed["chunk_count"], filename)
            print(
                f"PASS {filename}: document_id={doc_id}, "
                f"chunks={indexed['chunk_count']}, dim=512",
                flush=True,
            )

            if index == 0:
                # Indexing the same immutable document twice must not duplicate points.
                again = await client.post(f"/v1/documents/{doc_id}/index")
                again.raise_for_status()
                assert again.json()["chunk_count"] == indexed["chunk_count"]
                await assert_indexed(doc_id, indexed["chunk_count"], filename)
                print("PASS idempotent reindex", flush=True)

    print("M1b PASS", flush=True)


if __name__ == "__main__":
    asyncio.run(main())
