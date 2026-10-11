"""M1a smoke test: real TXT fixture and generated PDF/DOCX documents."""

from dataclasses import asdict
from io import BytesIO
from pathlib import Path
import json

import pymupdf
from docx import Document

from app.rag.parser import ParsedBlock, parse_document
from app.rag.chunker import chunk_document


def make_pdf() -> bytes:
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Synthetic contract: payment is due within 30 days.")
    data = pdf.tobytes()
    pdf.close()
    return data


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Payment Clause", level=1)
    doc.add_paragraph("Payment is due within 30 days after acceptance.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Party A"
    table.cell(0, 1).text = "Example Corporation"
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def main() -> None:
    cases = [
        ("sample_contract.txt", Path("data/samples/sample_contract.txt").read_bytes()),
        ("synthetic.pdf", make_pdf()),
        ("synthetic.docx", make_docx()),
    ]
    for filename, data in cases:
        blocks = parse_document(filename, data)
        chunks = chunk_document("demo-doc", filename, blocks)
        assert chunks and all(chunk.text for chunk in chunks)
        if filename.endswith(".pdf"):
            assert all(chunk.page == 1 for chunk in chunks)
        else:
            assert all(chunk.page is None for chunk in chunks)
        if filename.endswith(".docx"):
            assert any(chunk.section == "Payment Clause" for chunk in chunks)
            assert any("Example Corporation" in chunk.text for chunk in chunks)
        print(json.dumps(
            {"filename": filename, "blocks": len(blocks),
             "chunks": len(chunks), "first_chunk": asdict(chunks[0])},
            ensure_ascii=False,
        ))

    long_block = ParsedBlock(text="甲方应在验收后三十日内付款。" * 80, page=2)
    long_chunks = chunk_document(
        "long-demo", "long.pdf", [long_block], chunk_size=200, overlap=40
    )
    assert len(long_chunks) > 2
    assert all(c.page == 2 and len(c.text) <= 200 for c in long_chunks)
    assert len(set(c.chunk_id for c in long_chunks)) == len(long_chunks)
    print(f"Long-text split: {len(long_chunks)} chunks")
    print("M1a PASS")


if __name__ == "__main__":
    main()
