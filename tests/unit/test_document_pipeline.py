from io import BytesIO

import pymupdf
from docx import Document
import pytest

from app.rag.chunker import chunk_document
from app.rag.parser import ParsedBlock, parse_document


def test_txt_split_has_stable_ids() -> None:
    blocks = parse_document("contract.txt", ("甲方在验收后付款。" * 120).encode())
    first = chunk_document("doc-a", "contract.txt", blocks, chunk_size=200, overlap=40)
    again = chunk_document("doc-a", "contract.txt", blocks, chunk_size=200, overlap=40)
    assert len(first) >= 2
    assert [c.chunk_id for c in first] == [c.chunk_id for c in again]


def test_pdf_page_is_one_based() -> None:
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Payment due within 30 days.")
    data = pdf.tobytes()
    pdf.close()
    assert parse_document("contract.pdf", data)[0].page == 1


def test_docx_heading_and_table_are_preserved() -> None:
    doc = Document()
    doc.add_heading("Payment Clause", level=1)
    doc.add_paragraph("Pay within thirty days.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Amount"
    table.cell(0, 1).text = "10000"
    b = BytesIO()
    doc.save(b)
    blocks = parse_document("contract.docx", b.getvalue())
    assert blocks[0].section == "Payment Clause"
    assert blocks[0].page is None
    assert "10000" in blocks[0].text


def test_scanned_pdf_without_text_is_rejected() -> None:
    pdf = pymupdf.open()
    pdf.new_page()
    data = pdf.tobytes()
    pdf.close()
    with pytest.raises(ValueError, match="No extractable text"):
        parse_document("blank.pdf", data)


def test_chunks_never_span_pages() -> None:
    blocks = [
        ParsedBlock(text="第一页的付款约定。" * 60, page=1),
        ParsedBlock(text="第二页的违约责任。" * 60, page=2),
    ]
    chunks = chunk_document("doc-a", "contract.pdf", blocks, chunk_size=200, overlap=40)
    assert {c.page for c in chunks} == {1, 2}
    for chunk in chunks:
        assert ("第一页" in chunk.text) != ("第二页" in chunk.text)
